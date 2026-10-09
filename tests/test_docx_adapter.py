from docx import Document
from PIL import Image
from knowledge.processor.import_processor.nodes.docx_to_md_node import DocxToMdNode
from knowledge.utils.image_refs import image_references, resolve_image


def test_docx_preserves_order_table_heading_and_embedded_image(tmp_path):
    picture = tmp_path / "中文 图.png"
    Image.new("RGB", (20, 20), "green").save(picture)
    doc = Document()
    doc.add_heading("导入步骤", level=1)
    doc.add_paragraph("第一段正文")
    table = doc.add_table(rows=2, cols=2)
    for cell, value in zip(
        [c for r in table.rows for c in r.cells], ["项目", "限制", "文件", "40 MiB"]
    ):
        cell.text = value
    doc.add_picture(str(picture))
    doc.add_paragraph("图片后正文")
    doc.add_paragraph("# Python comment, not a heading")
    source = tmp_path / "中文 教学文档.docx"
    doc.save(source)
    result = DocxToMdNode().process(
        {"import_file_path": str(source), "file_dir": str(tmp_path)}
    )
    output = tmp_path / "docx" / "中文 教学文档.md"
    assert result["md_path"] == str(output)
    text = output.read_text(encoding="utf-8")
    assert (
        text.index("# 导入步骤")
        < text.index("| 项目 | 限制 |")
        < text.index("![文档内嵌图片]")
        < text.index("图片后正文")
    )
    refs = image_references(text)
    assert len(refs) == 1
    assert (
        resolve_image(output.parent, refs[0].destination).read_bytes()
        == picture.read_bytes()
    )
    assert "\\# Python comment, not a heading" in text
