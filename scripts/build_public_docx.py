"""Create the owned procedural test document, not a fabricated test result."""

from pathlib import Path
from docx import Document
from docx.shared import Inches

ROOT = Path(__file__).resolve().parents[1]
folder = ROOT / "examples/public/manual"
doc = Document()
for paragraph in (
    (folder / "掌柜智库 操作手册.md").read_text(encoding="utf-8").split("\n\n")
):
    if paragraph.startswith("#"):
        prefix, title = paragraph.split(" ", 1)
        doc.add_heading(title.strip(), level=min(len(prefix), 3))
    elif paragraph.startswith("!["):
        doc.add_picture(str(folder / "流程 图.png"), width=Inches(5.8))
    elif paragraph.strip():
        doc.add_paragraph(paragraph.strip())
doc.add_heading("格式选择表", level=2)
doc.add_paragraph("本表是操作规程，不是解析成功率或性能指标。")
table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
for cell, value in zip(
    table.rows[0].cells, ["格式", "推荐用途", "浏览器导入注意"], strict=True
):
    cell.text = value
for values in [
    ("PDF", "含图资料和固定版式", "使用 MinerU 解析"),
    ("DOCX", "可编辑的文字、内嵌图片与简单表格", "图片应嵌入文档"),
    ("Markdown", "纯文本知识资料", "单文件上传不携带相邻图片"),
]:
    for cell, value in zip(table.add_row().cells, values, strict=True):
        cell.text = value
target = folder / "掌柜智库 图文表格手册.docx"
doc.save(target)
print(target)
