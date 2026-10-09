"""DOCX adapter into the existing Markdown/image graph, without pretending it is PDF."""

import hashlib
import re
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from knowledge.processor.import_processor.base import BaseNode
from knowledge.processor.import_processor.state import ImportGraphState


class DocxToMdNode(BaseNode):
    name = "docx_to_md_node"

    def process(self, state: ImportGraphState) -> ImportGraphState:
        source = Path(state["import_file_path"])
        with ZipFile(source) as archive:
            if sum(i.file_size for i in archive.infolist()) > 256 * 1024 * 1024:
                raise ValueError("DOCX uncompressed content exceeds 256 MiB")
            if "word/document.xml" not in archive.namelist():
                raise ValueError("ZIP is not a Word document")
        document = Document(source)
        output = Path(state["file_dir"]) / "docx"
        images = output / "images"
        images.mkdir(parents=True, exist_ok=True)
        formats = {
            "image/png": ".png",
            "image/jpeg": ".jpg",
            "image/gif": ".gif",
            "image/bmp": ".bmp",
            "image/webp": ".webp",
        }

        def paragraph(block: Paragraph) -> str:
            text = block.text.rstrip()
            style = block.style.name if block.style else ""
            heading = re.fullmatch(r"(?:Heading|标题)\s*(\d+)", style, re.IGNORECASE)
            if heading and text:
                text = "#" * min(6, int(heading.group(1))) + " " + text
            elif text:
                # A literal Python comment in Word is not a Markdown heading.
                text = re.sub(r"(?m)^(\s{0,3})(#{1,6})(?=\s)", r"\1\\\2", text)
            for element in block._p.iter(qn("a:blip")):
                if element.get(qn("r:link")):
                    raise ValueError(
                        "Externally linked DOCX images are unsupported; embed the image first"
                    )
                rid = element.get(qn("r:embed"))
                if not rid:
                    raise ValueError("DOCX image has no embedded relationship")
                part = block.part.related_parts[rid]
                extension = formats.get(part.content_type)
                if not extension:
                    raise ValueError(
                        "Unsupported DOCX image type; export this document to PDF: "
                        + part.content_type
                    )
                digest = hashlib.sha256(part.blob).hexdigest()
                destination = images / (digest + extension)
                destination.write_bytes(part.blob)
                text += "\n\n![文档内嵌图片](images/" + destination.name + ")"
            # VML/OLE requires a renderer: do not silently lose its visual content.
            if any(
                e.tag.endswith("}imagedata") or e.tag == qn("w:object")
                for e in block._p.iter()
            ):
                raise ValueError(
                    "Legacy drawing/OLE found; export this document to PDF for faithful parsing"
                )
            return text

        pieces = []
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                pieces.append(paragraph(block))
            elif isinstance(block, Table):
                rows = []
                for row in block.rows:
                    cells = []
                    for cell in row.cells:
                        if cell.tables:
                            raise ValueError(
                                "Nested DOCX tables require PDF conversion"
                            )
                        cells.append(
                            "<br>".join(paragraph(p) for p in cell.paragraphs)
                            .replace("|", "\\|")
                            .replace("\n", "<br>")
                        )
                    rows.append("| " + " | ".join(cells) + " |")
                if rows:
                    pieces.append(
                        "\n".join(
                            [
                                rows[0],
                                "| "
                                + " | ".join("---" for _ in block.rows[0].cells)
                                + " |",
                                *rows[1:],
                            ]
                        )
                    )
        text = "\n\n".join(p for p in pieces if p.strip())
        if not text.strip():
            raise ValueError("DOCX has no extractable text or images")
        target = output / (source.stem + ".md")
        temporary = target.with_suffix(".md.tmp")
        temporary.write_text(text + "\n", encoding="utf-8")
        temporary.replace(target)
        return {"md_path": str(target)}
